"""Files an agent writes in any format the person opens (`create_file`): one registry entry per format, each turning the
same few fields — text, a table, sheets, a mail's header — into its bytes, locally and at no cost.

The formats and the content type each is served with are the contract's (`galaius_core.FILE_TYPES`); a format missing
here, or one here the contract does not name, fails at import."""

import csv
import io
import json
import os
import re
import stat
import subprocess
import sys
import tempfile
from abc import ABC, abstractmethod
from contextlib import suppress
from collections.abc import Iterable, Iterator
from email import message_from_bytes
from email.message import EmailMessage
from email.policy import default
from pathlib import Path, PurePosixPath
from typing import Annotated, ClassVar, Literal, NamedTuple

from docx import Document
from galaius_core import CREATE_FILE_TYPES, FILE_CONTENT_TYPES, file_key
from docx.shared import Mm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pptx import Presentation
from pydantic import AfterValidator, BaseModel, Field, model_validator

from galaius.agents import registry as reg
from galaius.agents.messaging import asking_run
from galaius.agents.providers import PROVIDERS
from galaius.agents.vocabulary import TouchScope
from galaius.caller import Caller
from galaius.fence import INSTRUCTIONS, STEERING
from galaius.pinned_directory import PinnedDirectory

#: One cell of a table: what a spreadsheet keeps as a value of its own type.
Cell = str | int | float | bool | None
Rows = list[list[Cell]]


def kept_path(path: str, suffixes: Iterable[str] | None = FILE_CONTENT_TYPES) -> str:
    """Relative to the working folder, '/'-separated and written plainly, never climbing, hidden, absolute or an agent's
    steering file (a CLAUDE.md written here would steer the next agent), with one of `suffixes` (every kept file's,
    `FILE_TYPES`, by default; None: any)."""
    allowed = None if suffixes is None else set(suffixes)
    parts = PurePosixPath(path).parts
    if (not parts or str(PurePosixPath(path)) != path or path.startswith(("/", "~")) or "\\" in path or ":" in path
            or any(part.startswith(".") or part.casefold() in STEERING | INSTRUCTIONS for part in parts)
            or (allowed is not None and PurePosixPath(path).suffix.lower() not in allowed)):
        raise ValueError(f"{path!r}: a path inside your working folder" + ("" if allowed is None else f" ending in one of {', '.join(sorted(allowed))}"))
    return path


def _address(address: str) -> str:
    """One plain address on one line: a line break would add a mail header of its own (a hidden Bcc)."""
    if not re.fullmatch(r"[^@\s<>,;\"]+@[^@\s<>,;\"]+\.[^@\s<>,;\"]+", address):
        raise ValueError(f"{address!r}: each recipient is one plain address (name@domain.tld)")
    return address


def _one_line(text: str) -> str:
    if re.search(r"[\r\n]", text):
        raise ValueError("a mail subject holds no line break")
    return text


#: A path `create_file` writes: one of its own formats (`CREATE_FILE_TYPES`).
PlainPath = Annotated[str, AfterValidator(lambda path: kept_path(path, dict(CREATE_FILE_TYPES)))]
Address = Annotated[str, AfterValidator(_address)]


class FileSpec(BaseModel):
    """What an agent asks `create_file` for. `path`'s suffix picks the format; each format reads the fields it takes."""

    path: PlainPath
    #: Text: the whole file (txt, md, json), the body (eml), Markdown-lite (docx, pptx: `#` headings, `-` bullets,
    #: paragraphs; a pptx slide per `#` heading), or Typst markup (pdf).
    content: str = ""
    #: A table: its first row is the header (csv, xlsx, docx, pdf).
    rows: Rows = Field(default_factory=list)
    #: Several named sheets (xlsx), each a table.
    sheets: dict[str, Rows] = Field(default_factory=dict)
    #: A mail draft's subject, recipients and attachments (eml); never sent. An attachment is a file this run wrote with
    #: create_file (`write` checks it), by its path.
    subject: Annotated[str, AfterValidator(_one_line)] = ""
    to: list[Address] = Field(default_factory=list)
    cc: list[Address] = Field(default_factory=list)
    attach: list[PlainPath] = Field(default_factory=list, max_length=10)
    #: The attachments' bytes by file name, as `write` read them (never given by an agent).
    attached: dict[str, bytes] = Field(default_factory=dict, exclude=True)

    @model_validator(mode="after")
    def attached_to_a_mail(self) -> "FileSpec":
        if self.attach and self.suffix != ".eml":
            raise ValueError("only a mail draft (.eml) takes attachments")
        return self

    @property
    def suffix(self) -> str:
        return PurePosixPath(self.path).suffix.lower()


class Block(NamedTuple):
    """One block of Markdown-lite text: a heading (its level), a bullet, a numbered item, or a paragraph (level 0)."""

    kind: Literal["heading", "bullet", "numbered", "paragraph"]
    text: str
    level: int = 0


class Span(NamedTuple):
    """A stretch of a block's text in one style: `**bold**`, `*italic*`, `` `code` `` or plain."""

    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


_INLINE = re.compile(r"\*\*\*(?P<both>.+?)\*\*\*|\*\*(?P<bold>.+?)\*\*|`(?P<code>[^`]+)`|(?<![\w*])\*(?P<italic>[^*\s][^*]*?)\*(?![\w*])")
#: The face a `code` span is set in, in a document and a deck alike.
CODE_FONT = "Consolas"


def spans(text: str) -> Iterator[Span]:
    """`text` cut into its styled stretches, the marks themselves dropped: a document shows bold, never `**`."""
    at = 0
    for found in _INLINE.finditer(text):
        if found.start() > at:
            yield Span(text[at:found.start()])
        yield Span(found["both"] or found["bold"] or found["code"] or found["italic"], bold=found["both"] is not None or found["bold"] is not None,
                   italic=found["both"] is not None or found["italic"] is not None, code=found["code"] is not None)
        at = found.end()
    if at < len(text):
        yield Span(text[at:])


def styled(font, span: Span) -> None:
    """`span`'s style on a run's `font` (a document's and a deck's take the same attributes)."""
    font.bold, font.italic = span.bold or None, span.italic or None
    if span.code:
        font.name = CODE_FONT


def plain(text: str) -> str:
    """`text` without its inline marks (a slide title, a sheet's words)."""
    return "".join(span.text for span in spans(text))


def blocks(text: str) -> Iterator[Block]:
    """`text` as headings (`#`..`######`), bullets (`-`, `*`), numbered items (`1.`, `1)`) and paragraphs (consecutive
    lines joined, a blank line between two)."""
    paragraph: list[str] = []
    for line in [*text.splitlines(), ""]:
        stripped = line.strip()
        heading = re.match(r"(#{1,6})\s+(.*)", stripped)
        bullet = re.match(r"[-*]\s+(.*)", stripped)
        numbered = re.match(r"\d{1,3}[.)]\s+(.*)", stripped)
        if paragraph and (not stripped or heading or bullet or numbered):
            yield Block("paragraph", " ".join(paragraph))
            paragraph = []
        if heading:
            yield Block("heading", heading[2], len(heading[1]))
        elif bullet:
            yield Block("bullet", bullet[1])
        elif numbered:
            yield Block("numbered", numbered[1])
        elif stripped:
            paragraph.append(stripped)


class FileFormat(ABC):
    """One format `create_file` writes: its suffixes and how a spec becomes its bytes. A leaf registers itself."""

    suffixes: ClassVar[tuple[str, ...]]
    #: The spec fields it needs filled, at least one of them.
    reads: ClassVar[tuple[str, ...]] = ("content",)
    registry: ClassVar[dict[str, "FileFormat"]] = {}

    def __init_subclass__(cls) -> None:
        super().__init_subclass__()
        for suffix in cls.suffixes:
            FileFormat.registry[suffix] = cls()

    @classmethod
    def of(cls, spec: FileSpec) -> "FileFormat":
        return cls.registry[spec.suffix]

    def bytes(self, spec: FileSpec) -> bytes:
        if not any(getattr(spec, field) for field in self.reads):
            raise ValueError(f"a {spec.suffix} file needs {' or '.join(self.reads)}")
        return self.render(spec)

    @abstractmethod
    def render(self, spec: FileSpec) -> bytes: ...


class Text(FileFormat):
    suffixes = (".txt", ".md")

    def render(self, spec: FileSpec) -> bytes:
        return spec.content.encode()


class Json(FileFormat):
    suffixes = (".json",)

    def render(self, spec: FileSpec) -> bytes:
        return json.dumps(json.loads(spec.content), indent=2, ensure_ascii=False).encode() + b"\n"


def _inert(cell: Cell) -> Cell:
    """A text cell a spreadsheet would run as a formula (`=`, `+`, `-`, `@`, a tab or a carriage return first) kept as
    text, as written; numbers stay numbers."""
    return "'" + cell if isinstance(cell, str) and cell[:1] in ("=", "+", "-", "@", "\t", "\r") else cell


class Csv(FileFormat):
    suffixes = (".csv",)
    reads = ("rows",)

    def render(self, spec: FileSpec) -> bytes:
        out = io.StringIO()
        csv.writer(out).writerows(["" if cell is None else _inert(cell) for cell in row] for row in spec.rows)
        return ("\ufeff" + out.getvalue()).encode()  # a byte-order mark: a spreadsheet reads it as UTF-8


class Eml(FileFormat):
    """A mail DRAFT: opened in a mail program, ready to check and send by the person; galaius never sends it."""

    suffixes = (".eml",)

    def render(self, spec: FileSpec) -> bytes:
        message = EmailMessage()
        for header, addresses in (("To", spec.to), ("Cc", spec.cc)):
            if addresses:
                message[header] = ", ".join(addresses)
        message["Subject"] = spec.subject
        message["X-Unsent"] = "1"  # opened as a draft to send, not a received mail
        message.set_content(spec.content)
        for name, data in spec.attached.items():
            kind, _, subtype = FILE_CONTENT_TYPES[PurePosixPath(name).suffix.lower()].partition("/")
            if kind == "message":  # a mail attached as itself (message/rfc822 takes no base64 encoding)
                message.add_attachment(message_from_bytes(data, policy=default), filename=name)
            else:
                message.add_attachment(data, maintype=kind, subtype=subtype, filename=name)
        return bytes(message)


class Docx(FileFormat):
    suffixes = (".docx",)
    reads = ("content", "rows")

    def render(self, spec: FileSpec) -> bytes:
        document = Document()
        for section in document.sections:  # A4, the page his printer and readers use (python-docx starts on Letter)
            section.page_width, section.page_height = Mm(210), Mm(297)
        for block in blocks(spec.content):
            paragraph = document.add_heading("", level=block.level) if block.kind == "heading" else \
                document.add_paragraph(style={"bullet": "List Bullet", "numbered": "List Number"}.get(block.kind))
            for span in spans(block.text):
                styled(paragraph.add_run(span.text).font, span)
        if spec.rows:
            width = max(len(row) for row in spec.rows)
            table = document.add_table(rows=0, cols=width, style="Table Grid")
            for number, row in enumerate(spec.rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = "" if value is None else str(value)
                    if number == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        out = io.BytesIO()
        document.save(out)
        return out.getvalue()


class Xlsx(FileFormat):
    """A workbook: a sheet per name, its first row the header (bold, kept in view, filterable), each column as wide as
    its longest cell up to `COLUMN_MAX` characters, a longer cell wrapped inside that width; printed on A4, one page wide."""

    suffixes = (".xlsx",)
    reads = ("rows", "sheets")
    COLUMN_MAX: ClassVar[int] = 60

    def render(self, spec: FileSpec) -> bytes:
        book = Workbook()
        book.remove(book.active)
        for name, rows in (spec.sheets or {"Sheet1": spec.rows}).items():
            sheet = book.create_sheet(_sheet_name(name))
            for row in rows:
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.data_type = "s"  # his text, never a formula the sheet runs
            if rows:
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(self.COLUMN_MAX, 2 + max(len(str(cell.value or "")) for cell in column))
                for cell in column:
                    if len(str(cell.value or "")) > self.COLUMN_MAX:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.page_setup.paperSize, sheet.page_setup.orientation = sheet.PAPERSIZE_A4, "landscape"
            sheet.sheet_properties.pageSetUpPr.fitToPage, sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = True, 1, 0
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()


def _sheet_name(name: str) -> str:
    """`name` as a sheet may be named: none of `[]:*?/\\`, at most 31 characters, never empty."""
    return re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Sheet"


class Pptx(FileFormat):
    """A deck: a slide per `#` heading, its bullets and paragraphs below it."""

    suffixes = (".pptx",)

    def render(self, spec: FileSpec) -> bytes:
        deck = Presentation()
        body = None
        for block in blocks(spec.content):
            if block.kind == "heading" or body is None:
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = plain(block.text) if block.kind == "heading" else ""
                body = slide.placeholders[1].text_frame
                body.text = ""
                if block.kind == "heading":
                    continue
            paragraph = body.paragraphs[0] if not body.paragraphs[0].runs else body.add_paragraph()
            for span in spans(block.text):
                run = paragraph.add_run()
                run.text = span.text
                styled(run.font, span)
        out = io.BytesIO()
        deck.save(out)
        return out.getvalue()


class Pdf(FileFormat):
    """Typst markup compiled here, a table after it when `rows` are given. Compiled apart (`COMPILE`): in a process of
    its own that bounds its own memory, stopped after `SECONDS`, from an empty folder (`#read` and `#image` reach no
    file of his), with no package and no network: every proxy points nowhere, which is what stops a package fetch
    (`@preview/…`); the `"@` check before only says so in plain words."""

    suffixes = (".pdf",)
    reads = ("content", "rows")
    #: The compiler's process, run with the source's file, its (empty) root folder and its memory bound in bytes. Exit 3:
    #: no compiler here (never the markup's fault).
    COMPILE: ClassVar[str] = (
        "import sys\n"
        "try:\n    import resource\n    resource.setrlimit(resource.RLIMIT_AS, (int(sys.argv[3]),) * 2)\nexcept ImportError:\n    pass\n"
        "try:\n    import typst\nexcept ImportError:\n    sys.exit(3)\n"
        "sys.stdout.buffer.write(typst.compile(sys.argv[1], root=sys.argv[2], package_path=sys.argv[2], package_cache_path=sys.argv[2]))\n")
    SECONDS: ClassVar[float] = 60.0
    MEMORY: ClassVar[int] = 2 * 1024**3
    #: The address every proxy variable names for the compiler: nothing listens there, so a package fetch fails.
    NOWHERE: ClassVar[str] = "http://127.0.0.1:9"

    def render(self, spec: FileSpec) -> bytes:
        source = spec.content
        if re.search(r'"@', source):
            raise ValueError("a PDF here imports no Typst package (`@preview/...`): write the markup itself")
        if spec.rows:
            width = max(len(row) for row in spec.rows)
            cells = ", ".join(_typst_string(cell) for row in spec.rows for cell in [*row, *[None] * (width - len(row))])
            source += f"\n\n#table(columns: {width}, {cells})\n"
        offline = {name: self.NOWHERE for name in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "http_proxy", "https_proxy", "all_proxy")}
        with tempfile.TemporaryDirectory(prefix="galaius-typst-") as root:
            main = Path(root) / "main.typ"
            main.write_text(source, encoding="utf-8")
            try:
                done = subprocess.run([sys.executable, "-I", "-c", self.COMPILE, str(main), root, str(self.MEMORY)], capture_output=True,
                                      timeout=self.SECONDS, env={**os.environ, **offline, "NO_PROXY": "", "no_proxy": ""})
            except subprocess.TimeoutExpired:
                raise ValueError(f"the PDF did not compile within {self.SECONDS:.0f} s") from None
        if done.returncode == 3:
            raise OSError("this computer's galaius has no PDF compiler (typst): reinstall galaius")
        if done.returncode or not done.stdout.startswith(b"%PDF"):
            said = done.stderr.decode(errors="replace").strip().splitlines()
            raise ValueError("the Typst markup did not compile: " + (said[-1] if said else f"exit {done.returncode}"))
        return done.stdout


def _typst_string(value: Cell) -> str:
    """`value` as a Typst string literal: quoted, its backslashes and quotes escaped, control characters dropped."""
    text = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", "", "" if value is None else str(value))
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n") + '"'


if FileFormat.registry.keys() != dict(CREATE_FILE_TYPES).keys():
    raise RuntimeError(f"create_file formats {sorted(FileFormat.registry)} differ from the contract's {sorted(dict(CREATE_FILE_TYPES))}")


class Written(NamedTuple):
    path: Path
    size: int
    #: Its kept copy's key (`reg.RunFiles`) when an agent run asked; None for a person's own session.
    key: str | None


def write(spec: FileSpec, *, replace: bool) -> Written:
    """`spec` written as its format in the caller's working folder (`save`), a mail draft's attachments read first
    (`_attachments`: for an agent run, only files it wrote itself)."""
    if spec.attach:
        spec = spec.model_copy(update={"attached": _attachments(spec.attach, _asking())})
    return save(spec.path, FileFormat.of(spec).bytes(spec), replace=replace)


def _asking() -> "reg.AgentRun | None":
    """The agent run asking (`asking_run`: its process ancestry first, never only what its environment claims); None for a
    person's own session. PermissionError: this process names a run this computer does not know."""
    asked = asking_run()
    if asked is None and Caller.current().environment.get("GALAIUS_RUN_ID"):
        raise PermissionError("this process names an agent run this computer does not know")
    return reg.get_run(asked) if asked is not None else None


def writable(path: str) -> tuple[str, Caller, "reg.AgentRun | None", Path]:
    """Where `save` would write `path`, refused (ValueError, PermissionError) where it may not: a path of `kept_path`, a
    run that reads only, a fenced run's level that is no `write` / `sandbox`. Asked before a paid job (`galaius.media`),
    so nothing is charged that cannot be kept."""
    path = kept_path(path)
    caller = Caller.current()
    run = _asking()
    target = caller.cwd.resolve() / path
    if run is not None and (provider := PROVIDERS.get(run.provider)) is not None and provider.touch_of(run.permission_mode) == TouchScope.READ_ONLY:
        raise PermissionError("this run reads only: it writes no file")
    if run is not None and run.fence is not None:
        _reachable(run.fence, target, WRITE_LEVELS, "write")
    return path, caller, run, target


def save(path: str, data: bytes, *, replace: bool) -> Written:
    """`data` saved at `path` (`kept_path`: one of `FILE_TYPES`) in the caller's working folder: only where the caller's
    run may write (its fence's `write` / `sandbox` levels, its project folder; anywhere in that folder for a run with no
    fence, which writes there itself), every folder on the way reached without a link and made when missing, a file
    already there only when `replace`, written whole or not at all; then kept for the web (`reg.RunFiles`) under the run
    asking (`_asking`). What `create_file` and every media tool write (`galaius.media`)."""
    path, caller, run, target = writable(path)
    *folders, name = PurePosixPath(path).parts
    with PinnedDirectory.open(caller.cwd.resolve(), *folders, create=True) as folder:
        staged = f".{name}.{os.urandom(6).hex()}.new"
        with os.fdopen(folder.file(staged, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o644), "wb") as handle:
            handle.write(data)
        try:
            if replace:
                folder.replace(staged, folder, name)
            else:
                folder.link(staged, folder, name)  # fails when the name is taken, even by a file made a moment ago
        except FileExistsError:
            raise ValueError(f"{path} exists already: pass replace to write over it") from None
        finally:
            with suppress(FileNotFoundError):
                folder.unlink(staged)
    return Written(target, len(data), reg.RunFiles.keep(run.run_id, data, PurePosixPath(path).suffix.lower()) if run is not None else None)


#: The most a mail draft's attachments weigh together: the draft stays under what its chat card downloads
#: (`MACHINE_AGENT_MEDIA`, once base64-encoded inside the mail).
ATTACHED_MAX = 4 * 1024 * 1024


def _attachments(paths: list[str], run: "reg.AgentRun | None") -> dict[str, bytes]:
    """The files a mail draft attaches, by name, read from the working folder (`read_working`). For an agent run, only a file it wrote
    itself with create_file (its kept copy, `reg.RunFiles`, holds the same bytes): the draft is downloadable from the
    web, and must never carry out a file of his it merely found in the folder (an `.env`, his taxes)."""
    found: dict[str, bytes] = {}
    for path in paths:
        name = PurePosixPath(path).name
        data = read_working(path, ATTACHED_MAX)
        if name in found:
            raise ValueError(f"two attachments are named {name}")
        if run is not None and not reg.RunFiles.path(run.run_id, file_key(data, PurePosixPath(name).suffix.lower())).is_file():
            raise PermissionError(f"{path}: a draft attaches only a file you wrote with create_file in this run")
        found[name] = data
    if sum(map(len, found.values())) > ATTACHED_MAX:
        raise ValueError(f"the attachments weigh more than {ATTACHED_MAX // (1024 * 1024)} MiB together")
    return found


#: The levels at which a fenced run writes a file of its working folder (`save`), and reads one (`read_working`).
WRITE_LEVELS: frozenset[str] = frozenset({"write", "sandbox"})
READ_LEVELS: frozenset[str] = frozenset({"read", "write_on_review", "sandbox", "write"})


def read_working(path: str, max_bytes: int) -> bytes:
    """File `path` of the caller's working folder, read whole: a plain path below it (`kept_path`, any suffix), reached
    without a link, a plain file, at most `max_bytes`; for a fenced run only where its levels let agents read
    (`READ_LEVELS`, its project folder). ValueError / PermissionError: why not. What a mail draft attaches and what a
    media job reads (`galaius.media`)."""
    path = kept_path(path, None)
    root = Caller.current().cwd.resolve()
    run = _asking()
    if run is not None and run.fence is not None:
        _reachable(run.fence, root / path, READ_LEVELS, "read")
    *folders, name = PurePosixPath(path).parts
    try:
        with PinnedDirectory.open(root, *folders) as folder, os.fdopen(folder.file(name, os.O_RDONLY | PinnedDirectory.NONBLOCK), "rb") as handle:
            if not stat.S_ISREG(os.fstat(handle.fileno()).st_mode):
                raise ValueError(f"{path} is not a plain file")
            data = handle.read(max_bytes + 1)
    except FileNotFoundError:
        raise ValueError(f"{path}: no such file in your working folder") from None
    if len(data) > max_bytes:
        raise ValueError(f"{path} weighs more than the {max_bytes / (1024 * 1024):g} MiB it may")
    return data


def _reachable(fence, target: Path, levels: frozenset[str], verb: str) -> None:
    """Refuse `target` unless the fenced run may `verb` it: one of `levels`, or its project folder."""
    places = fence.places()
    if fence.project is not None and target.is_relative_to(fence.project):
        return
    try:
        parts = target.relative_to(places.base).parts
    except ValueError:
        raise PermissionError(f"{target} is outside this PC's working directory") from None
    level = places.reach(parts)
    if level not in levels:
        raise PermissionError(f"{'/'.join(parts[:-1]) or '.'} is at level « {level} » on this PC: agents {verb} only where it is « {' » or « '.join(sorted(levels))} »")
